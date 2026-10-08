import os
import re
import uuid
from datetime import datetime
import requests
from flask import Flask, request, jsonify, send_from_directory
from flask_cors import CORS
from docx import Document
from openpyxl import Workbook

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 10 * 1024 * 1024
CORS(app)

# === CẤU HÌNH ===
def env(name, default=""):
    val = os.environ.get(name, default)
    return val.strip() if isinstance(val, str) else val

GEMINI_API_KEY = env("GEMINI_API_KEY")
GEMINI_MODEL = env("GEMINI_MODEL", "gemini-2.0-flash")
GROQ_API_KEY = env("GROQ_API_KEY")
GROQ_MODEL = env("GROQ_MODEL", "llama-3.3-70b-versatile")

TIMEOUT = 120
OUTPUT_FOLDER = "outputs"
os.makedirs(OUTPUT_FOLDER, exist_ok=True)
DEFAULT_PROMPT = "Trả lời bằng tiếng Việt rõ ràng, chính xác, phù hợp lĩnh vực thủy điện."

# === GỌI AI ===
def call_gemini(prompt, system=DEFAULT_PROMPT):
    if not GEMINI_API_KEY:
        return None, "Chưa có GEMINI_API_KEY"
    try:
        url = f"https://generativelanguage.googleapis.com/v1/models/{GEMINI_MODEL}:generateContent?key={GEMINI_API_KEY}"
        full = f"Hướng dẫn: {system}\nNội dung: {prompt}"
        r = requests.post(url, json={"contents": [{"parts": [{"text": full}]}]}, timeout=TIMEOUT)
        r.raise_for_status()
        return r.json()["candidates"][0]["content"]["parts"][0]["text"].strip(), None
    except Exception as e:
        return None, str(e)

def call_groq(messages):
    if not GROQ_API_KEY:
        return None, "Chưa có GROQ_API_KEY"
    try:
        r = requests.post(
            "https://api.groq.com/openai/v1/chat/completions",
            headers={"Authorization": f"Bearer {GROQ_API_KEY}", "Content-Type": "application/json"},
            json={"model": GROQ_MODEL, "messages": messages, "max_tokens": 2048, "temperature": 0.7},
            timeout=TIMEOUT
        )
        r.raise_for_status()
        return r.json()["choices"][0]["message"]["content"].strip(), None
    except Exception as e:
        return None, str(e)

def chat_ai(prompt, system=DEFAULT_PROMPT):
    msg = [{"role": "system", "content": system}, {"role": "user", "content": prompt}]
    for name, fn in [("Gemini", lambda: call_gemini(prompt, system)), ("Groq", lambda: call_groq(msg))]:
        text, err = fn()
        if text:
            return f"✅ [{name}]\n{text}", text
    return "❌ Kiểm tra lại API key", ""

# === XUẤT FILE ===
def clean(t): return re.sub(r"[*_`#]", "", t)

def make_word(content):
    try:
        doc = Document()
        doc.add_heading("BÁO CÁO", 0)
        doc.add_paragraph(f"Ngày: {datetime.now().strftime('%d/%m/%Y %H:%M')}")
        doc.add_paragraph("")
        for p in content.split("\n"):
            if p.strip(): doc.add_paragraph(clean(p))
        fn = f"baocao_{uuid.uuid4().hex[:6]}.docx"
        doc.save(os.path.join(OUTPUT_FOLDER, fn))
        return fn
    except: return None

def make_excel(content):
    try:
        wb = Workbook()
        ws = wb.active
        ws.append([f"Ngày: {datetime.now().strftime('%d/%m/%Y %H:%M')}"])
        ws.append(["Nội dung"])
        for line in content.split("\n"):
            if line.strip(): ws.append([clean(line)])
        fn = f"baocao_{uuid.uuid4().hex[:6]}.xlsx"
        wb.save(os.path.join(OUTPUT_FOLDER, fn))
        return fn
    except: return None

# === ROUTES ===
@app.route("/")
def index():
    return '''<!DOCTYPE html>
<html lang="vi">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>All Thủy Điện</title>
<style>
*{margin:0;padding:0;box-sizing:border-box;font-family:Segoe UI,sans-serif}
:root{--p:#2563eb;--s:#16a34a;--g:#f59e0b}
body{background:linear-gradient(135deg,#eff6ff,#f0fdf4);min-height:100vh;padding:20px}
h1{text-align:center;color:var(--p);margin-bottom:24px;font-size:22px}
.grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(360px,1fr));gap:20px;max-width:1300px;margin:0 auto}
.card{background:#fff;border-radius:16px;padding:20px;box-shadow:0 4px 12px rgba(0,0,0,.06);display:flex;flex-direction:column;height:650px}
.card h2{font-size:16px;color:var(--p);margin-bottom:12px;padding-bottom:8px;border-bottom:2px solid #e2e8f0}
.chat{flex:1;overflow-y:auto;padding:12px;background:#f8fafc;border-radius:12px;margin-bottom:12px}
.msg{margin:8px 0;max-width:90%}
.msg.u{margin-left:auto}
.msg.a{margin-right:auto}
.bubble{padding:12px 16px;border-radius:16px;line-height:1.5;font-size:14px}
.u .bubble{background:var(--p);color:#fff;border-bottom-right-radius:4px}
.a .bubble{background:#e2e8f0;color:#1e293b;border-bottom-left-radius:4px}
textarea{width:100%;padding:12px;border:1px solid #e2e8f0;border-radius:12px;font-size:14px;resize:none;height:60px;margin-bottom:8px}
button{padding:10px 16px;border:none;border-radius:10px;font-size:14px;font-weight:600;cursor:pointer}
.b1{background:var(--p);color:#fff}
.b2{background:var(--s);color:#fff}
.b3{background:var(--g);color:#fff}
.btn-row{display:flex;gap:8px}
.info{font-size:13px;line-height:2;color:#475569}
hr{border:none;border-top:1px solid #e2e8f0;margin:16px 0}
</style>
</head>
<body>
<h1>⚡ All Thủy Điện — Hệ thống hỗ trợ</h1>
<div class="grid">
  <!-- Cột 1 -->
  <div class="card">
    <h2>✍️ Soạn thảo văn bản</h2>
    <div class="chat" id="c1"></div>
    <textarea id="i1" placeholder="Nhập yêu cầu..."></textarea>
    <div class="btn-row">
      <button class="b1" onclick="go(1)">Gửi</button>
      <button class="b2" onclick="dl(1,'word')">Word</button>
      <button class="b3" onclick="dl(1,'excel')">Excel</button>
    </div>
  </div>

  <!-- Cột 2 -->
  <div class="card">
    <h2>📊 Quy trình & Báo cáo</h2>
    <div class="chat" id="c2"></div>
    <textarea id="i2" placeholder="Nhập yêu cầu..."></textarea>
    <div class="btn-row">
      <button class="b1" onclick="go(2)">Gửi</button>
      <button class="b2" onclick="dl(2,'word')">Word</button>
      <button class="b3" onclick="dl(2,'excel')">Excel</button>
    </div>
  </div>

  <!-- Cột 3 -->
  <div class="card">
    <h2>📌 Thông tin liên hệ</h2>
    <div class="info" style="flex:1">
      <p><strong>Admin:</strong> Chotvjp</p>
      <p><strong>ĐT:</strong> 0973020486</p>
      <p><strong>Gmail:</strong> hoangdien86ncc@gmail.com</p>
      <p><strong>Bộ phận:</strong> Quản lý kỹ thuật</p>
      <p><strong>Đơn vị:</strong> Công ty CP Thủy điện Nậm Chiến</p>
      <p><strong>Địa chỉ:</strong> TK5 - Mường La - Sơn La</p>
      <p><strong>Website:</strong> namchien.vn</p>
      <hr>
      <p style="color:#94a3b8">© 2026 — Hệ thống hỗ trợ AI</p>
    </div>
  </div>
</div>

<script>
let last = {1:"",2:""};
const sys = {1:"Bạn soạn thảo văn bản chính thức, chuẩn tiếng Việt trang trọng.",
             2:"Bạn là chuyên gia thủy điện, phân tích dữ liệu, lập quy trình, báo cáo kỹ thuật chi tiết."};

function add(col,role,t){
  const box = document.getElementById(`c${col}`);
  const d = document.createElement("div");
  d.className = `msg ${role}`;
  d.innerHTML = `<div class="bubble">${t.replace(/\n/g,"<br>")}</div>`;
  box.appendChild(d);
  box.scrollTop = box.scrollHeight;
  if(role==="a") last[col]=t;
}

async function go(col){
  const ta = document.getElementById(`i${col}`);
  const txt = ta.value.trim();
  if(!txt) return;
  add(col,"u",txt); ta.value="";
  add(col,"a","⏳ Đang xử lý...");
  const res = await fetch("/api/chat",{
    method:"POST", headers:{"Content-Type":"application/json"},
    body:JSON.stringify({p:txt,s:sys[col]})
  });
  const d = await res.json();
  const box = document.getElementById(`c${col}`);
  box.lastChild.remove();
  add(col,"a",d.rpy||"Lỗi");
}

async function dl(col,type){
  if(!last[col]) return alert("Chưa có nội dung để tải");
  const res = await fetch("/api/export",{
    method:"POST", headers:{"Content-Type":"application/json"},
    body:JSON.stringify({txt:last[col],type:type})
  });
  const d = await res.json();
  if(d.url) window.open(d.url,"_blank");
  else alert("Lỗi: "+(d.err||""));
}
</script>
</body>
</html>
'''

@app.route("/api/chat", methods=["POST"])
def api_chat():
    d = request.get_json(silent=True) or {}
    rpy, raw = chat_ai(d.get("p",""), d.get("s", DEFAULT_PROMPT))
    return jsonify({"rpy": rpy, "raw": raw})

@app.route("/api/export", methods=["POST"])
def api_export():
    d = request.get_json(silent=True) or {}
    txt = d.get("txt","")
    if not txt: return jsonify({"err": "Không có nội dung"}), 400
    fn = make_word(txt) if d.get("type")=="word" else make_excel(txt)
    if not fn: return jsonify({"err": "Lỗi tạo file"}), 500
    return jsonify({"url": f"/get/{fn}"})

@app.route("/get/<fn>")
def get_file(fn):
    return send_from_directory(OUTPUT_FOLDER, fn, as_attachment=True)

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 10000)))
